#!/usr/bin/env python3
"""Batpilot hackathon pitch deck — brand-matched, fully editable .pptx."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

HERE = os.path.dirname(os.path.abspath(__file__))
PAPER = RGBColor(0xFA, 0xF8, 0xEF)
INK = RGBColor(0x17, 0x15, 0x0C)
MUTED = RGBColor(0x6F, 0x6A, 0x54)
ACID = RGBColor(0xE3, 0xE6, 0x16)
GREEN = RGBColor(0x3E, 0x7A, 0x34)
RED = RGBColor(0xBE, 0x4A, 0x2F)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def bg(slide, color=PAPER):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))


def para(tf, text, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Arial"
    return p


def title_slide(slide, kicker, title, sub):
    bg(slide)
    tb = textbox(slide, 1, 0.6, 11.3, 1)
    para(tb.text_frame, kicker, 16, True, MUTED, first=True)
    tb = textbox(slide, 1, 1.6, 11.3, 2.6)
    para(tb.text_frame, title, 54, True, first=True)
    tb = textbox(slide, 1, 4.4, 11.3, 1.4)
    para(tb.text_frame, sub, 22, False, MUTED, first=True)


def section(slide, num, title, bullets, accent_last=False):
    bg(slide)
    tb = textbox(slide, 1, 0.5, 11.3, 1)
    para(tb.text_frame, num, 16, True, MUTED, first=True)
    tb = textbox(slide, 1, 1.2, 11.3, 1.2)
    para(tb.text_frame, title, 40, True, first=True)
    tb = textbox(slide, 1, 2.8, 11.3, 4.2)
    for i, b in enumerate(bullets):
        first = i == 0
        c = GREEN if (accent_last and i == len(bullets) - 1) else INK
        para(tb.text_frame, "\u2022  " + b, 20, False, c, first=first)
        if not first:
            tb.text_frame.paragraphs[-1].space_before = Pt(14)


def image_slide(slide, num, title, img, caption):
    bg(slide)
    tb = textbox(slide, 1, 0.5, 11.3, 1)
    para(tb.text_frame, num, 16, True, MUTED, first=True)
    tb = textbox(slide, 1, 1.1, 11.3, 1)
    para(tb.text_frame, title, 36, True, first=True)
    slide.shapes.add_picture(img, Inches(1), Inches(2.3), width=Inches(11.3))
    tb = textbox(slide, 1, 6.7, 11.3, 0.6)
    para(tb.text_frame, caption, 14, False, MUTED, first=True)


# 1 cover
s = prs.slides.add_slide(blank)
title_slide(s, "ARBITRUM OPEN HOUSE SINGAPORE · ROBINHOOD CHAIN",
            "Batpilot",
            "Your US stocks, managed while you sleep.\nRecurring buys + stop-loss protection on tokenized stocks — settled in USDG, verifiable onchain.\ngithub.com/PhiBao/batpilot  ·  batpilot.vercel.app")

# 2 problem
s = prs.slides.add_slide(blank)
section(s, "01 · PROBLEM", "Asia sleeps through every US session",
        ["US market hours are 21:30–04:00 in Singapore — holders sleep through every session.",
         "Tokenized stocks trade 24/7 but are priced 24/5: feeds go stale on weekends.",
         "No onchain DCA exists; only perps have stop-losses. Spot holders babysit broker apps at 3am."])

# 3 insight
s = prs.slides.add_slide(blank)
section(s, "02 · INSIGHT", "Treat market sessions as onchain state",
        ["A guard screens every fill: freshness, volatility-adaptive bands, corporate actions.",
         "Refusals emit events — the fill it refuses is the proof it is safe.",
         "After any protection sale, the plan cools down instead of buying back the crash."],
        accent_last=True)

# 4 product
s = prs.slides.add_slide(blank)
section(s, "03 · PRODUCT", "Set once. Covered around the clock.",
        ["Fund a plan: $50 NVDA every 15 min, −8% stop, +20% take, 2% slip, 2h cooldown.",
         "Permissionless keeper executes; contracts decide. Idle USDG earns in Steakhouse Earn.",
         "Every fill, refusal and protection links to its transaction."])

# 5 live proof
s = prs.slides.add_slide(blank)
section(s, "04 · LIVE PROOF", "Real money, reconciled to the cent",
        ["Mainnet plan filled by keeper: $5 NVDA @ $228.82 — tx 0x12e29b94…db5689.",
         "Equity decoded word-by-word: $34 cash + 0.0665 NVDA + $14.99 value. Exact.",
         "Testnet loop filling on 15-min cadence, fully verified on Blockscout."],
        accent_last=True)

# 6 tech
s = prs.slides.add_slide(blank)
section(s, "05 · TECHNICAL DEPTH", "A real engine, honestly benchmarked",
        ["Volatility-adaptive bands from Chainlink round history — no trusted inputs.",
         "Ported to Stylus Rust, deployed onchain: bit-identical outputs at 41% less gas.",
         "31 Foundry tests + 7 Stylus tests, incl. mainnet-fork proofs."])

# 7 architecture
s = prs.slides.add_slide(blank)
image_slide(s, "06 · ARCHITECTURE", "Four boxes and a trail",
            os.path.join(HERE, "arch.png"),
            "Vault + guard onchain · Uniswap v3 + Steakhouse Earn venues · permissionless keeper · everything verifiable.")

# 8 app
s = prs.slides.add_slide(blank)
image_slide(s, "07 · PRODUCT", "Designed like an editorial, not a dashboard",
            os.path.join(HERE, "..", "screenshot.png"),
            "batpilot.vercel.app — chain switcher, living plan sentences, sparklines, infinite verifiable trail.")

# 9 gtm
s = prs.slides.add_slide(blank)
section(s, "08 · GO-TO-MARKET", "From sleep-deprived holders to family offices",
        ["Wedge: SG/MY retail holding US stocks — concierge onboarding, track-record cards.",
         "Expand: autopilot baskets, SGD rail, advisor white-label.",
         "Money: freemium fills + yield take-rate, aligned with Global Dollar Network."])

# 10 roadmap/ask
s = prs.slides.add_slide(blank)
section(s, "09 · ROADMAP", "Milestone-ready",
        ["Funded mainnet plan + public track record → oraclePaused hardening → TWAMM + baskets.",
         "Multisig + audit before scaling external funds.",
         "Ask: buildathon prize funds the audit; Founder House turns it into family-office rails."])

# 11 close
s = prs.slides.add_slide(blank)
title_slide(s, "BATPILOT",
            "Arcus gives you a venue. Batpilot gives you sleep.",
            "github.com/PhiBao/batpilot  ·  batpilot.vercel.app  ·  Set it once. Verified onchain.")

out = os.path.join(HERE, "Batpilot_Pitch.pptx")
prs.save(out)
print("saved", out)
